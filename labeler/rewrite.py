"""改写台核心逻辑：建议定稿标签 + 改写记录 IO + 最终改写导出。

双人复核没跑完也能开始出改写：每条笔记用「已有的最好标签」当建议定稿——
双人一致 > 单人复核 > 机器（两人分歧回退机器预填，页面上再人工定夺）。
分桶语义与 adjudicate.compare_pair / review_tool.merge_reviews 完全一致。

改写记录按操作人分文件 data/rewrites/{操作人}.json；LLM 草稿集中存
data/rewrites/drafts.jsonl（单条按钮和 rewrite_batch 共用，带 input_sig 过期标记）。
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import openpyxl

from labeler.export_xlsx import _new_wb
from labeler.notes import Note
from labeler.review_tool import (
    ReviewRecord,
    _fmt_ids,
    machine_intent_ids,
    machine_persona_lines,
    machine_scen_ids,
)
from labeler.taxonomy import Taxonomy

REWRITE_HEADER = [
    "note_id", "标题", "原正文(完整)", "改写query", "改写context", "状态",
    "最终场景(路径)", "最终意图", "需补充信息", "备注", "依据", "操作人", "时间",
]

STATUS_CN = {"keep": "保留", "drop": "不需要", "insufficient": "缺少信息"}
MAX_REWRITE_LEN = 200
MAX_CONTEXT_LEN = 2000   # context=原帖清洗版（仅作兜底截断，正常不该触到）

# ── context 清洗（纯代码，不走 LLM）─────────────────────────────────
_TAG_RE = re.compile(r"#[^#\n]{1,30}#?")                 # #话题# / #话题
_EMOTION_CODE_RE = re.compile(r"\[[^\]\n]{1,10}R\]")     # [哭惹R] 等表情码
_AT_RE = re.compile(r"@[^\s@，。,；;：:]{1,30}")           # @提及
_EDIT_MARK_RE = re.compile(r"(?:^|\n)\s*[—\-–=]{0,4}\s*[二三四五]编(?=\s|[:：])")  # 二编：起为追加段
_EMOJI_RE = re.compile("[\U0001F000-\U0001FAFF\u2600-\u27BF\u2B00-\u2BFF\uFE0F\u200d]")


def clean_context(note: Note) -> str:
    """原帖清洗版 context：保留原文措辞、语气和全部实质信息，不概括不改述；
    只删 话题标签/表情码/emoji/@提及，砍掉「二编/三编」起的追加段，压缩连排标点。"""
    title = (note.title or "").strip()
    desc = (note.desc or "").strip()
    text = desc if not title or desc.startswith(title) else f"{title}\n{desc}"
    m = _EDIT_MARK_RE.search(text)
    if m:                                             # 二编/三编 是后补内容，整段砍掉
        text = text[:m.start()]
    text = _TAG_RE.sub("", text)
    text = _EMOTION_CODE_RE.sub("", text)
    text = _EMOJI_RE.sub("", text)
    text = _AT_RE.sub("", text)
    text = re.sub(r"([!！?？])\1+", r"\1", text)
    lines = [ln.strip() for ln in text.splitlines()]
    squeezed: list[str] = []
    for ln in lines:
        if ln or (squeezed and squeezed[-1]):          # 连续空行压成一行
            squeezed.append(ln)
    text = "\n".join(squeezed).strip()
    return text[:MAX_CONTEXT_LEN] + ("…" if len(text) > MAX_CONTEXT_LEN else "")


@dataclass
class Suggestion:
    """一条笔记的建议定稿标签（分桶 + 标签 + 来源）。"""

    bucket: str        # agree 双人一致 / drop 双剔除 / conflict 分歧 / single 仅单人 / machine 仅机器
    based_on: str      # 建议依据：如 "WH+zlx" / "WH" / "机器(两人分歧)"
    final_status: str  # keep | drop | insufficient
    scenarios: list[str] = field(default_factory=list)
    intents: list[str] = field(default_factory=list)
    persona: list[str] = field(default_factory=list)
    missing_info: str = ""


@dataclass
class RewriteRecord:
    """单条改写台定稿：最终标签 + 改写 query（会话 JSON 与导出共用）。"""

    note_id: str
    final_status: str = ""  # keep | drop | insufficient（空 = 未定稿）
    scenarios: list[str] = field(default_factory=list)
    intents: list[str] = field(default_factory=list)
    persona: list[str] = field(default_factory=list)
    missing_info: str = ""
    comment: str = ""
    rewrite: str = ""
    context: str = ""  # 原帖里对回答 query 有用的背景（rewrite_v2 起生成，可空）
    draft: str = ""  # 最近一次 AI 草稿（留档，草稿明细在 drafts.jsonl）
    based_on: str = ""
    operator: str = ""
    updated_at: str = ""


def suggest_final(cmp_: dict[str, Any], ra: ReviewRecord | None, rb: ReviewRecord | None,
                  rec: dict[str, Any], left: str, right: str) -> Suggestion:
    """compare_pair 结论 → 建议定稿标签。agree 取左序（集合已相等），single 取该方，
    drop/none 各自显然，conflict 预填机器标签等人工定夺。"""
    kind = cmp_["kind"]
    if kind == "agree":
        assert ra is not None and rb is not None
        va, vb = cmp_["verdict_a"], cmp_["verdict_b"]
        miss = "；".join(filter(None, [
            ra.missing_info if va == "insufficient" else "",
            rb.missing_info if vb == "insufficient" else ""]))
        return Suggestion(
            bucket="agree", based_on=f"{left}+{right}",
            final_status="insufficient" if "insufficient" in (va, vb) else "keep",
            scenarios=list(ra.scenarios), intents=list(ra.intents),
            persona=list(ra.persona), missing_info=miss)
    if kind == "drop":
        return Suggestion(bucket="drop", based_on=f"{left}+{right}", final_status="drop")
    if kind == "conflict":
        return Suggestion(
            bucket="conflict", based_on="机器(两人分歧)", final_status="keep",
            scenarios=machine_scen_ids(rec), intents=machine_intent_ids(rec),
            persona=machine_persona_lines(rec))
    if kind in ("single_a", "single_b"):
        rv = ra if kind == "single_a" else rb
        v = cmp_["verdict_a"] if kind == "single_a" else cmp_["verdict_b"]
        assert rv is not None and v is not None
        return Suggestion(
            bucket="single", based_on=left if kind == "single_a" else right,
            final_status=v, scenarios=list(rv.scenarios), intents=list(rv.intents),
            persona=list(rv.persona), missing_info=rv.missing_info)
    assert kind == "none"
    return Suggestion(
        bucket="machine", based_on="机器", final_status="keep",
        scenarios=machine_scen_ids(rec), intents=machine_intent_ids(rec),
        persona=machine_persona_lines(rec))


def rewrite_to_dict(r: RewriteRecord) -> dict[str, Any]:
    return {
        "note_id": r.note_id, "final_status": r.final_status,
        "scenarios": r.scenarios, "intents": r.intents, "persona": r.persona,
        "missing_info": r.missing_info, "comment": r.comment,
        "rewrite": r.rewrite, "context": r.context, "draft": r.draft,
        "based_on": r.based_on,
        "operator": r.operator, "updated_at": r.updated_at,
    }


def rewrite_from_dict(d: dict[str, Any]) -> RewriteRecord:
    return RewriteRecord(
        note_id=str(d.get("note_id") or ""),
        final_status=str(d.get("final_status") or ""),
        scenarios=[str(s) for s in d.get("scenarios") or []],
        intents=[str(i) for i in d.get("intents") or []],
        persona=[str(p) for p in d.get("persona") or []],
        missing_info=str(d.get("missing_info") or ""),
        comment=str(d.get("comment") or ""),
        rewrite=str(d.get("rewrite") or ""),
        context=str(d.get("context") or ""),
        draft=str(d.get("draft") or ""),
        based_on=str(d.get("based_on") or ""),
        operator=str(d.get("operator") or ""),
        updated_at=str(d.get("updated_at") or ""),
    )


def load_operator_session(path: Path) -> dict[str, dict[str, Any]]:
    if not path.exists():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    return data if isinstance(data, dict) else {}


def load_drafts(path: Path) -> dict[str, dict[str, Any]]:
    """草稿 jsonl → {note_id: 最新草稿记录}（后写覆盖；无 draft 的失败记录跳过）。"""
    drafts: dict[str, dict[str, Any]] = {}
    if not path.exists():
        return drafts
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        try:
            d = json.loads(line)
        except json.JSONDecodeError:
            continue
        nid = str(d.get("note_id") or "")
        if nid and d.get("draft"):
            drafts[nid] = d
    return drafts


def append_draft(path: Path, record: dict[str, Any]) -> None:
    """追加一条草稿记录（网页单条按钮与 rewrite_batch 批量共用）。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as f:
        f.write(json.dumps(record, ensure_ascii=False) + "\n")


def input_sig(final_status: str, scenarios: list[str], intents: list[str]) -> str:
    """草稿过期判定键：定稿分诊/场景/意图变了，旧草稿即失效。"""
    return f"{final_status}|{','.join(scenarios)}|{','.join(intents)}"


_FENCE_RE = re.compile(r"^```[\w-]*\s*|\s*```$")
_PREFIX_RE = re.compile(r"^(改写\s*query|query|改写)\s*[:：]\s*", re.IGNORECASE)
_QUOTES = (("「", "」"), ("『", "』"), ("“", "”"), ('"', '"'), ("'", "'"))


def clean_rewrite_output(raw: str) -> str:
    """LLM 草稿清洗：剥 ``` 围栏/「改写query：」前缀/包裹引号，折叠空白，截断 200 字。"""
    text = (raw or "").strip()
    if text.startswith("```"):
        text = _FENCE_RE.sub("", text).strip()
    text = _PREFIX_RE.sub("", text).strip()
    for lq, rq in _QUOTES:
        if len(text) >= 2 and text.startswith(lq) and text.endswith(rq):
            text = text[1:-1].strip()
    text = re.sub(r"\s+", " ", text)
    return text[:MAX_REWRITE_LEN] + ("…" if len(text) > MAX_REWRITE_LEN else "")


def split_context_query(raw: str) -> tuple[str, str]:
    """rewrite_v2 的 LLM 输出 → (context, query)。

    期望 {"context": ..., "query": ...} JSON（允许围栏/前后杂文字）；解析不出
    有效 query 时整段当 query（context 置空，与 rewrite_v1 行为一致）。
    """
    text = (raw or "").strip()
    if text.startswith("```"):
        text = _FENCE_RE.sub("", text).strip()
    m = re.search(r"\{.*\}", text, re.DOTALL)
    if m:
        try:
            d = json.loads(m.group(0))
        except json.JSONDecodeError:
            d = None
        if isinstance(d, dict):
            context = re.sub(r"\s+", " ", str(d.get("context") or "")).strip()
            context = context[:MAX_CONTEXT_LEN] + ("…" if len(context) > MAX_CONTEXT_LEN else "")
            query = clean_rewrite_output(str(d.get("query") or ""))
            if query:
                return context, query
    return "", clean_rewrite_output(raw)


def export_rewrites(task: list[tuple[Note, dict[str, Any]]], tax: Taxonomy,
                    records: dict[str, dict[str, Any]], out_path: Path,
                    ) -> tuple[Path, dict[str, int]]:
    """改写会话 → 最终改写 xlsx（最终改写 / 待改写 / 统计）。

    完成判定与页面进度一致：drop，或（定稿且改写非空）。
    """
    final_rows: list[list[Any]] = []
    todo_rows: list[list[Any]] = []
    n_done = n_drop = n_rewritten = 0

    for note, _rec in task:
        nid = note.note_id
        ar = records.get(nid)
        if ar is None:
            todo_rows.append([nid, note.title, "未处理", ""])
            continue
        status = str(ar.get("final_status") or "")
        rewrite = str(ar.get("rewrite") or "").strip()
        if status not in STATUS_CN:
            todo_rows.append([nid, note.title, "分诊未定", rewrite])
            continue
        if status != "drop" and not rewrite:
            todo_rows.append([nid, note.title, "保留但未改写 query", ""])
            continue
        n_done += 1
        if status == "drop":
            n_drop += 1
        else:
            n_rewritten += 1
        final_rows.append([
            nid, note.title, note.desc, rewrite,
            str(ar.get("context") or ""), STATUS_CN[status],
            _fmt_ids([str(s) for s in ar.get("scenarios") or []], tax, "scenario"),
            _fmt_ids([str(i) for i in ar.get("intents") or []], tax, "intent"),
            str(ar.get("missing_info") or ""), str(ar.get("comment") or ""),
            str(ar.get("based_on") or ""), str(ar.get("operator") or ""),
            str(ar.get("updated_at") or ""),
        ])

    wb = _new_wb()
    ws = wb.create_sheet("最终改写")
    ws.append(REWRITE_HEADER)
    for row in final_rows:
        ws.append(row)
    for col, width in zip("ABCDEFGHIJKLM",
                          (24, 28, 50, 44, 40, 22, 50, 10, 24, 24, 14, 10, 20)):
        ws.column_dimensions[col].width = width
    ws2 = wb.create_sheet("待改写")
    ws2.append(["note_id", "标题", "情况", "改写query"])
    for row in todo_rows:
        ws2.append(row)
    for col, width in zip("ABCD", (24, 28, 22, 50)):
        ws2.column_dimensions[col].width = width
    stats: list[list[Any]] = [
        ["任务总数", len(task)],
        ["已完成（剔除或已改写）", n_done],
        ["其中剔除（不需要）", n_drop],
        ["已改写 query", n_rewritten],
        ["待改写条数", len(todo_rows)],
        ["导出时间", datetime.now(UTC).isoformat(timespec="seconds")],
    ]
    ws3 = wb.create_sheet("统计")
    ws3.append(["指标", "数值"])
    for row in stats:
        ws3.append(row)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path, {"total": len(task), "done": n_done, "drop": n_drop,
                      "rewritten": n_rewritten, "todo": len(todo_rows)}


def open_workbook_check(path: Path) -> None:  # pragma: no cover - 供外部快速校验
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        assert wb.sheetnames == ["最终改写", "待改写", "统计"]
    finally:
        wb.close()
