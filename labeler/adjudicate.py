"""裁决工具：两位复核人结果的逐条比对 + 最终改写。

流程（在 review_server 的 /adjudicate 页面操作）::

    1. 选左/右两位复核人（对方交回的 Excel 先用复核页「导入Excel」落到其名下会话）
    2. 逐条比对 A / B / 机器 三方结论，快捷「以A/以B/以机器为准」再微调
    3. 每条保留的笔记在改写框里把噪声原帖改写成标准 query
    4. 导出 最终标注_*.xlsx（最终结果 / 待处理 / 统计）

比对语义与 review_tool.merge_reviews 完全一致：
双 drop=剔除；双方可采纳（保留/缺少信息）且场景意图一致=一致；其余=待裁决。
裁决记录按 (左, 右) 一对存一份 data/adjudications/{左}__{右}.json。
"""

from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import openpyxl

from labeler.export_xlsx import _new_wb
from labeler.notes import Note
from labeler.review_tool import ReviewRecord, _fmt_ids, _verdict

FINAL_HEADER = [
    "note_id", "标题", "原正文(完整)", "改写query", "分诊",
    "最终场景", "最终意图", "需补充信息", "备注", "裁决人", "裁决时间",
]


def compare_pair(ra: ReviewRecord | None, rb: ReviewRecord | None,
                 rec: dict[str, Any]) -> dict[str, Any]:
    """A/B 两条复核记录的比对结论（与 merge_reviews 的分桶规则一致）。"""
    va, vb = _verdict(ra, rec), _verdict(rb, rec)
    if va and vb:
        assert ra is not None and rb is not None
        if va == vb == "drop":
            kind = "drop"                      # 双不需要：一致剔除
        elif (va != "drop" and vb != "drop"
              and set(ra.scenarios) == set(rb.scenarios)
              and set(ra.intents) == set(rb.intents)):
            kind = "agree"                     # 标签一致：直接采纳
        else:
            kind = "conflict"                  # 待裁决
    elif va:
        kind = "single_a"
    elif vb:
        kind = "single_b"
    else:
        kind = "none"
    return {"kind": kind, "verdict_a": va, "verdict_b": vb}


_KIND_CN = {"drop": "双不需要", "agree": "标签一致", "conflict": "两人不一致",
            "single_a": "仅A完成", "single_b": "仅B完成", "none": "两人均未完成"}


def export_adjudications(
    task: list[tuple[Note, dict[str, Any]]], tax: Any,
    records: dict[str, dict[str, Any]], out_path: Path,
) -> tuple[Path, dict[str, int]]:
    """裁决会话 → 最终标注 xlsx（最终结果 / 待处理 / 统计）。"""
    final_rows: list[list[Any]] = []
    todo_rows: list[list[Any]] = []
    n_decided = n_drop = n_rewritten = 0

    for note, rec in task:
        nid = note.note_id
        ar = records.get(nid)
        if ar is None:
            todo_rows.append([nid, note.title, "未裁决", ""])
            continue
        status = str(ar.get("final_status") or "")
        rewrite = str(ar.get("rewrite") or "").strip()
        if status not in ("keep", "drop", "insufficient"):
            todo_rows.append([nid, note.title, "分诊未定", rewrite])
            continue
        if status != "drop" and not rewrite:
            todo_rows.append([nid, note.title, "保留但未改写 query", ""])
        n_decided += 1
        if status == "drop":
            n_drop += 1
        elif rewrite:
            n_rewritten += 1
        final_rows.append([
            nid, note.title, note.desc, rewrite,
            {"keep": "保留", "drop": "不需要", "insufficient": "缺少信息"}[status],
            _fmt_ids([str(s) for s in ar.get("scenarios") or []], tax, "scenario"),
            _fmt_ids([str(i) for i in ar.get("intents") or []], tax, "intent"),
            str(ar.get("missing_info") or ""), str(ar.get("comment") or ""),
            str(ar.get("adjudicator") or ""), str(ar.get("decided_at") or ""),
        ])

    wb = _new_wb()
    ws = wb.create_sheet("最终结果")
    ws.append(FINAL_HEADER)
    for row in final_rows:
        ws.append(row)
    for col, width in zip("ABCDEFGHIJK", (24, 28, 50, 50, 10, 40, 20, 24, 24, 10, 20)):
        ws.column_dimensions[col].width = width
    ws2 = wb.create_sheet("待处理")
    ws2.append(["note_id", "标题", "情况", "改写query"])
    for row in todo_rows:
        ws2.append(row)
    for col, width in zip("ABCD", (24, 28, 22, 50)):
        ws2.column_dimensions[col].width = width
    stats: list[list[Any]] = [
        ["任务总数", len(task)],
        ["已裁决", n_decided],
        ["其中剔除（不需要）", n_drop],
        ["保留且已改写 query", n_rewritten],
        ["待处理条数", len(todo_rows)],
        ["导出时间", datetime.now(UTC).isoformat(timespec="seconds")],
    ]
    ws3 = wb.create_sheet("统计")
    ws3.append(["指标", "数值"])
    for row in stats:
        ws3.append(row)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path, {"total": len(task), "decided": n_decided, "drop": n_drop,
                      "rewritten": n_rewritten, "todo": len(todo_rows)}


def load_pair_session(path: Path) -> dict[str, dict[str, Any]]:
    if not path.exists():
        return {}
    import json

    data = json.loads(path.read_text(encoding="utf-8"))
    return data if isinstance(data, dict) else {}


def open_workbook_check(path: Path) -> None:  # pragma: no cover - 供外部快速校验
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        assert wb.sheetnames == ["最终结果", "待处理", "统计"]
    finally:
        wb.close()
