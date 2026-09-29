"""export_xlsx 测试：总表结构、零覆盖叶保留、缺标条目入审核页签、分表齐全。"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl

from labeler.export_xlsx import export_all


def _rec(nid: str, scenarios: list[dict], intents: list[str],
         persona: list[dict], unlabelable: dict | None = None) -> dict:
    return {
        "note_id": nid, "status": "ok",
        "annotation": {
            "scenarios": scenarios, "intents": intents, "persona": persona,
            "unlabelable": unlabelable or {"scenario": False, "intent": False, "reason": ""},
            "warnings": [],
        },
    }


def _write_notes(path: Path) -> None:
    long_desc = "很长的正文" * 120  # >400 字符，验证不截断
    rows = [
        {"note_id": "n1", "title": "考研还是就业", "desc": long_desc, "tags": []},
        {"note_id": "n2", "title": "求安慰", "desc": "延毕了", "tags": []},
        {"note_id": "n3", "title": "缺标的那条", "desc": "反复失败", "tags": []},
    ]
    path.write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in rows), encoding="utf-8")


def test_export_all(tmp_path: Path, taxonomy) -> None:
    notes_path = tmp_path / "notes.jsonl"
    _write_notes(notes_path)
    recs = [
        _rec("n1", [{"id": "S001", "confidence": 0.9}, {"id": "S003", "confidence": 0.4}],
             ["I01", "I02"],
             [{"field": "年级", "value": "大三", "basis": "明确", "evidence": "我大三"}]),
        _rec("n2", [], [],
             [{"field": "年龄段", "value": "22左右", "basis": "推测", "evidence": ""}],
             {"scenario": True, "intent": True, "reason": "纯情绪宣泄"}),
    ]
    out_dir = tmp_path / "exports"
    written = export_all(recs, notes_path, taxonomy, "testrun", out_dir)

    master = out_dir / "标注总表_testrun.xlsx"
    assert master in written
    l1_dir = out_dir / "一级场景"
    for l1 in ("升学", "就业", "个人成长与生活"):
        assert (l1_dir / f"{l1}.xlsx") in written

    wb = openpyxl.load_workbook(master)
    assert set(wb.sheetnames) == {
        "全部标注", "场景汇总", "意图汇总", "画像汇总", "未完成与待审核",
    }

    # 场景汇总：全部叶子都在，零覆盖的 S004 计数为 0（空着让人看见）
    rows = list(wb["场景汇总"].iter_rows(min_row=2, values_only=True))
    by_sid = {r[0]: r for r in rows}
    assert set(by_sid) == {"S001", "S002", "S003", "S004"}
    assert by_sid["S001"][3] == 1 and by_sid["S004"][3] == 0

    # 全部标注：无场景无意图的 n2（无法标注）不收录，只留 n1；正文完整不截断
    master_rows = list(wb["全部标注"].iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in master_rows] == ["n1"]
    assert master_rows[0][2] == "很长的正文" * 120

    # 明细：多场景笔记在两个一级分表里各出现一行；正文同样完整
    wb升学 = openpyxl.load_workbook(l1_dir / "升学.xlsx")
    detail = list(wb升学["明细"].iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in detail] == ["n1"]
    assert detail[0][2] == "很长的正文" * 120

    # 画像汇总：只算高质量子集，n2（仅画像、无法标注）的推测不再计入
    rows = list(wb["画像汇总"].iter_rows(min_row=2, values_only=True))
    stat = {r[1]: (r[3], r[4]) for r in rows}
    assert stat["年级"] == (1, 0) and stat["年龄段"] == (0, 0)

    # 未完成与待审核：缺标 n3 在列，n2 的无法标注也在列
    rows = list(wb["未完成与待审核"].iter_rows(min_row=2, values_only=True))
    kinds = {r[0]: r[2] for r in rows}
    assert kinds["n3"] == "缺标(反复失败)"
    assert kinds["n2"] == "场景、意图无法标注"
