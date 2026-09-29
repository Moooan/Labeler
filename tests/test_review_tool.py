"""review_tool 测试：任务构建/预填、导出→导入往返（含清空语义）、双人合并一致性。"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl

from labeler.review_tool import (
    ReviewRecord,
    _is_reviewed,
    build_task,
    export_workbook,
    import_workbook,
    merge_reviews,
)


def _setup(tmp_path: Path) -> tuple[Path, list]:
    notes_path = tmp_path / "notes.jsonl"
    rows = [
        {"note_id": "n1", "title": "考研还是就业", "desc": "大三很纠结", "tags": []},
        {"note_id": "n2", "title": "只求安慰", "desc": "挂科了睡不着", "tags": []},
        {"note_id": "n3", "title": "低质量", "desc": "美食探店", "tags": []},
    ]
    notes_path.write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in rows),
                          encoding="utf-8")
    annotations = [
        {"note_id": "n1", "status": "ok", "annotation": {
            "scenarios": [{"id": "S001", "confidence": 0.9}, {"id": "S003", "confidence": 0.4}],
            "intents": ["I01", "I02"],
            "persona": [{"field": "年级", "value": "大三", "basis": "明确", "evidence": "我大三"}]}},
        {"note_id": "n2", "status": "ok", "annotation": {
            "scenarios": [], "intents": ["I03"], "persona": []}},
        {"note_id": "n3", "status": "ok", "annotation": {
            "scenarios": [], "intents": [], "persona": [
                {"field": "年龄段", "value": "20+", "basis": "推测", "evidence": ""}]}},
    ]
    return notes_path, annotations


def test_build_task_filters_low_quality(tmp_path: Path, taxonomy) -> None:
    notes_path, annotations = _setup(tmp_path)
    task = build_task(annotations, notes_path)
    assert [n.note_id for n, _ in task] == ["n1", "n2"]  # n3 无场景无意图，默认不进任务
    task_all = build_task(annotations, notes_path, include_unlabeled=True)
    assert len(task_all) == 3


def test_export_prefill_and_roundtrip(tmp_path: Path, taxonomy) -> None:
    notes_path, annotations = _setup(tmp_path)
    task = build_task(annotations, notes_path)
    out = tmp_path / "task_A.xlsx"
    export_workbook(task, taxonomy, {}, "A", out)

    wb = openpyxl.load_workbook(out)
    ws = wb["复核"]
    rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert set(rows) == {"n1", "n2"}
    assert rows["n1"][4] == "S001 S003"      # 复核场景 预填机器结果
    assert rows["n1"][8] == "I01 I02"        # 复核意图
    assert rows["n1"][13] in ("", None)      # 状态：未复核（空单元格读回 None）
    assert rows["n1"][11] in ("", None)      # 需补充信息：未分诊为空
    assert set(wb.sheetnames) == {"复核", "场景速查", "意图速查"}
    assert wb["场景速查"].max_row == 5        # 表头 + 4 叶

    # 复核人 A 改了结论（S002），n2 判定无场景；导出→导入 往返一致
    records = {
        "n1": ReviewRecord(note_id="n1", scenarios=["S002"], intents=["I01"],
                           persona=["年级: 大四 [明确]"], comment="只留一个", status="keep",
                           missing_info="", reviewer="A", reviewed_at="2026-09-24"),
        "n2": ReviewRecord(note_id="n2", scenarios=[], intents=["I03"], status="keep",
                           reviewer="A", reviewed_at="2026-09-24"),
    }
    out2 = tmp_path / "task_A_done.xlsx"
    export_workbook(task, taxonomy, records, "A", out2)
    rows2 = {r[0]: r for r in openpyxl.load_workbook(out2)["复核"]
             .iter_rows(min_row=2, values_only=True)}
    assert rows2["n1"][13] == "保留"
    back = import_workbook(out2, taxonomy)
    assert back["n1"].scenarios == ["S002"] and back["n1"].comment == "只留一个"
    assert back["n1"].persona == ["年级: 大四 [明确]"]
    assert back["n2"].scenarios == [] and back["n2"].status == "keep"  # 清空不被回填
    # 分诊三态 + 需补充信息 的中文列 ↔ 内部状态 往返
    tri = {"n1": ReviewRecord(note_id="n1", status="drop", reviewer="A"),
           "n2": ReviewRecord(note_id="n2", status="insufficient", scenarios=["S004"],
                              intents=["I03"], missing_info="目标专业不明", reviewer="A")}
    out3 = tmp_path / "task_A_triage.xlsx"
    export_workbook(task, taxonomy, tri, "A", out3)
    rows3 = {r[0]: r for r in openpyxl.load_workbook(out3)["复核"]
             .iter_rows(min_row=2, values_only=True)}
    assert rows3["n1"][13] == "不需要" and rows3["n2"][13] == "缺少信息"
    assert rows3["n2"][11] == "目标专业不明"
    back3 = import_workbook(out3, taxonomy)
    assert back3["n1"].status == "drop" and back3["n2"].status == "insufficient"
    assert back3["n2"].missing_info == "目标专业不明"
    # 未动过的行（预填=机器结果、无状态）不算已复核；手改了标签但没写状态算已复核
    untouched = import_workbook(out, taxonomy)
    assert not _is_reviewed(untouched["n1"], annotations[0])
    hand = ReviewRecord(note_id="n1", scenarios=["S002"], intents=["I01", "I02"])
    assert _is_reviewed(hand, annotations[0])


def test_merge_agreement_and_conflict(tmp_path: Path, taxonomy) -> None:
    notes_path, annotations = _setup(tmp_path)
    task = build_task(annotations, notes_path)
    a = {
        "n1": ReviewRecord(note_id="n1", scenarios=["S002"], intents=["I01"], status="done",
                           persona=["年级: 大四 [明确]"], reviewer="A"),
        "n2": ReviewRecord(note_id="n2", scenarios=[], intents=["I03"], status="done",
                           bad_fit=True, new_scenario="个人成长与生活 > 情绪 > 失眠焦虑",
                           reviewer="A"),
    }
    b = {
        "n1": ReviewRecord(note_id="n1", scenarios=["S002"], intents=["I01"], status="done",
                           persona=["年级: 大三 [明确]"], reviewer="B"),
        "n2": ReviewRecord(note_id="n2", scenarios=["S004"], intents=["I03"], status="done",
                           reviewer="B"),
    }
    out = tmp_path / "merged.xlsx"
    path, st = merge_reviews(a, b, task, taxonomy, out)
    assert path == out
    assert st["both"] == 2 and st["agree"] == 1 and st["conflict"] == 1

    wb = openpyxl.load_workbook(out)
    assert wb["采纳一致"].max_row == 2                       # n1 一致采纳
    agree_row = next(wb["采纳一致"].iter_rows(min_row=2, values_only=True))
    assert agree_row[0] == "n1"
    assert "升学 > 考研 > 院校选择" in agree_row[2]
    conflict = list(wb["不一致待裁决"].iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in conflict] == ["n2"]                # n2 不一致单独保存
    assert conflict[0][3] == "（无）" and "拖延" in conflict[0][4]
    sugg = list(wb["新场景建议"].iter_rows(min_row=2, values_only=True))
    assert len(sugg) == 1 and sugg[0][2] == "A" and "失眠焦虑" in sugg[0][4]
    assert wb["待另一人复核"].max_row == 1                    # 没有单人完成


def test_merge_triage_verdicts(tmp_path: Path, taxonomy) -> None:
    """先分诊后打标：双不需要→一致剔除；双缺少信息且标签一致→采纳（带需补充）；
    分诊或标签不一致→待裁决。"""
    notes_path, annotations = _setup(tmp_path)
    task = build_task(annotations, notes_path)
    out = tmp_path / "merged_triage.xlsx"

    a = {
        "n1": ReviewRecord(note_id="n1", status="drop", comment="广告水贴", reviewer="A"),
        "n2": ReviewRecord(note_id="n2", status="insufficient", reviewer="A"),  # 没打标
    }
    b = {
        "n1": ReviewRecord(note_id="n1", status="drop", reviewer="B"),
        "n2": ReviewRecord(note_id="n2", scenarios=["S004"], intents=["I03"],
                           status="keep", reviewer="B"),
    }
    _, st = merge_reviews(a, b, task, taxonomy, out)
    assert st["drop"] == 1 and st["insufficient"] == 0 and st["conflict"] == 1

    wb = openpyxl.load_workbook(out)
    drops = list(wb["一致剔除"].iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in drops] == ["n1"]
    assert drops[0][2] == "不需要" and drops[0][3] == "广告水贴"
    conflict = list(wb["不一致待裁决"].iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in conflict] == ["n2"]
    assert "缺少信息" in conflict[0][3]              # A 侧标签+缺少信息标记
    assert "拖延" in conflict[0][4]                  # B 侧照常显示标签
    stats = {r[0]: r[1] for r in wb["统计"].iter_rows(min_row=2, values_only=True)}
    assert stats["一致（剔除·不需要）"] == 1

    # 双缺少信息、标签一致 → 采纳一致（缺少信息=是，需补充两人合并）
    out2 = tmp_path / "merged_ins.xlsx"
    _, st2 = merge_reviews(
        {"n2": ReviewRecord(note_id="n2", status="insufficient", scenarios=["S004"],
                            intents=["I03"], missing_info="目标专业", reviewer="A")},
        {"n2": ReviewRecord(note_id="n2", status="insufficient", scenarios=["S004"],
                            intents=["I03"], missing_info="当前年级", reviewer="B")},
        task, taxonomy, out2)
    assert st2["insufficient"] == 1 and st2["agree"] == 0 and st2["conflict"] == 0
    row2 = next(openpyxl.load_workbook(out2)["采纳一致"].iter_rows(min_row=2, values_only=True))
    assert row2[6] == "是" and "目标专业" in row2[7] and "当前年级" in row2[7]
