"""adjudicate 测试：A/B 比对分类（与 merge 分桶一致）+ 最终标注导出。"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from labeler.adjudicate import compare_pair, export_adjudications
from labeler.review_tool import ReviewRecord


def _setup(tmp_path: Path) -> tuple[Path, list]:
    notes_path = tmp_path / "notes.jsonl"
    import json

    rows = [
        {"note_id": "n1", "title": "考研还是就业", "desc": "大三很纠结", "tags": []},
        {"note_id": "n2", "title": "只求安慰", "desc": "挂科了睡不着", "tags": []},
    ]
    notes_path.write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in rows),
                          encoding="utf-8")
    annotations = [
        {"note_id": "n1", "status": "ok", "annotation": {
            "scenarios": [{"id": "S001", "confidence": 0.9}], "intents": ["I01"], "persona": []}},
        {"note_id": "n2", "status": "ok", "annotation": {
            "scenarios": [], "intents": ["I03"], "persona": []}},
    ]
    return notes_path, annotations


def test_compare_pair_kinds(taxonomy) -> None:
    rec = {"note_id": "n1", "annotation": {"scenarios": [{"id": "S001"}], "intents": ["I01"]}}
    keep = ReviewRecord(note_id="n1", scenarios=["S001"], intents=["I01"], status="keep")
    other = ReviewRecord(note_id="n1", scenarios=["S002"], intents=["I01"], status="keep")
    drop = ReviewRecord(note_id="n1", status="drop")
    ins = ReviewRecord(note_id="n1", scenarios=["S001"], intents=["I01"],
                       status="insufficient", missing_info="年级")

    assert compare_pair(keep, keep, rec)["kind"] == "agree"
    assert compare_pair(drop, drop, rec)["kind"] == "drop"
    assert compare_pair(keep, other, rec)["kind"] == "conflict"
    assert compare_pair(keep, drop, rec)["kind"] == "conflict"
    assert compare_pair(ins, keep, rec)["kind"] == "agree"      # 缺少信息与保留标签一致=一致
    assert compare_pair(keep, None, rec)["kind"] == "single_a"
    assert compare_pair(None, None, rec)["kind"] == "none"
    assert compare_pair(None, drop, rec)["kind"] == "single_b"


def test_export_adjudications(tmp_path: Path, taxonomy) -> None:
    notes_path, annotations = _setup(tmp_path)
    from labeler.review_tool import build_task

    task = build_task(annotations, notes_path)
    records = {
        "n1": {"final_status": "keep", "scenarios": ["S001"], "intents": ["I01"],
               "persona": ["年级: 大三 [明确]"], "missing_info": "", "comment": "ok",
               "rewrite": "大三学生纠结考研还是就业，想要决策建议", "based_on": "a",
               "adjudicator": "C", "decided_at": "2026-09-24"},
        "n2": {"final_status": "drop", "scenarios": [], "intents": [],
               "rewrite": "", "based_on": "a", "adjudicator": "C", "decided_at": "2026-09-24"},
    }
    out = tmp_path / "最终标注.xlsx"
    path, st = export_adjudications(task, taxonomy, records, out)
    assert path == out
    assert st == {"total": 2, "decided": 2, "drop": 1, "rewritten": 1, "todo": 0}

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["最终结果", "待处理", "统计"]
    rows = {r[0]: r for r in wb["最终结果"].iter_rows(min_row=2, values_only=True)}
    assert rows["n1"][3] == "大三学生纠结考研还是就业，想要决策建议"   # 改写query列
    assert rows["n1"][4] == "保留"
    assert rows["n2"][4] == "不需要"
    assert wb["待处理"].max_row == 1                                     # 都处理完了

    # 保留但没改写 → 进待处理，不算完成
    records["n1"]["rewrite"] = ""
    _, st2 = export_adjudications(task, taxonomy, records, out)
    assert st2["rewritten"] == 0 and st2["todo"] == 1
