"""改写台测试：建议定稿分桶 + 草稿清洗/IO + 最终改写导出。"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from labeler.adjudicate import compare_pair
from labeler.notes import Note
from labeler.prompts import RewriteInput, build_rewrite_prompt
from labeler.review_tool import ReviewRecord, build_task
from labeler.rewrite import (
    RewriteRecord,
    clean_context,
    clean_rewrite_output,
    export_rewrites,
    load_drafts,
    rewrite_from_dict,
    rewrite_to_dict,
    split_context_query,
    suggest_final,
)


def _setup(tmp_path: Path) -> tuple[Path, list]:
    notes_path = tmp_path / "notes.jsonl"
    rows = [
        {"note_id": "n1", "title": "考研还是就业", "desc": "大三很纠结", "tags": []},
        {"note_id": "n2", "title": "只求安慰", "desc": "挂科了睡不着", "tags": []},
    ]
    notes_path.write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in rows),
                          encoding="utf-8")
    annotations = [
        {"note_id": "n1", "status": "ok", "annotation": {
            "scenarios": [{"id": "S001", "confidence": 0.9}], "intents": ["I01"], "persona": []}},
        {"note_id": "n2", "status": "ok", "annotation": {
            "scenarios": [], "intents": ["I03"], "persona": []}},
    ]
    return notes_path, annotations


def test_suggest_final_buckets(taxonomy) -> None:
    rec = {"note_id": "n1", "annotation": {
        "scenarios": [{"id": "S001"}], "intents": ["I01"], "persona": []}}
    keep = ReviewRecord(note_id="n1", scenarios=["S001"], intents=["I01"], status="keep",
                        persona=["年级: 大三 [明确]"])
    other = ReviewRecord(note_id="n1", scenarios=["S002"], intents=["I01"], status="keep")
    drop = ReviewRecord(note_id="n1", status="drop")
    ins = ReviewRecord(note_id="n1", scenarios=["S001"], intents=["I01"],
                       status="insufficient", missing_info="年级")

    # 双人一致：取左序（keep 在左带画像），insufficient 传染并合并双方需补充
    s = suggest_final(compare_pair(keep, ins, rec), keep, ins, rec, "WH", "zlx")
    assert s.bucket == "agree" and s.based_on == "WH+zlx"
    assert s.final_status == "insufficient" and s.missing_info == "年级"
    assert s.scenarios == ["S001"] and s.persona == ["年级: 大三 [明确]"]

    # 双剔除
    s = suggest_final(compare_pair(drop, drop, rec), drop, drop, rec, "WH", "zlx")
    assert s.bucket == "drop" and s.final_status == "drop" and not s.scenarios

    # 分歧 → 机器预填，等人工定夺
    s = suggest_final(compare_pair(keep, other, rec), keep, other, rec, "WH", "zlx")
    assert s.bucket == "conflict" and s.based_on == "机器(两人分歧)"
    assert s.scenarios == ["S001"] and s.final_status == "keep"

    # 对方 skip 退化为仅单人
    skipper = ReviewRecord(note_id="n1", status="skip")
    cmp_ = compare_pair(keep, skipper, rec)
    assert cmp_["kind"] == "single_a"
    s = suggest_final(cmp_, keep, skipper, rec, "WH", "zlx")
    assert s.bucket == "single" and s.based_on == "WH" and s.final_status == "keep"

    # 仅右且其结论为 drop
    s = suggest_final(compare_pair(None, drop, rec), None, drop, rec, "WH", "zlx")
    assert s.bucket == "single" and s.based_on == "zlx" and s.final_status == "drop"

    # 两人均未复核 → 机器
    s = suggest_final(compare_pair(None, None, rec), None, None, rec, "WH", "zlx")
    assert s.bucket == "machine" and s.based_on == "机器"
    assert s.scenarios == ["S001"] and s.intents == ["I01"]


def test_clean_rewrite_output() -> None:
    assert clean_rewrite_output("  大三纠结考研还是就业，求建议  ") == "大三纠结考研还是就业，求建议"
    assert clean_rewrite_output("```\n\"query here\"\n```") == "query here"
    assert clean_rewrite_output("改写query：我是问句") == "我是问句"
    assert clean_rewrite_output("Query: english prefix") == "english prefix"
    assert clean_rewrite_output("「围起来的问句」") == "围起来的问句"
    assert clean_rewrite_output("a\n\nb   c") == "a b c"
    assert len(clean_rewrite_output("字" * 250)) == 201        # 截断 200 + 省略号
    assert clean_rewrite_output("") == ""


def test_load_drafts_last_wins(tmp_path: Path) -> None:
    p = tmp_path / "drafts.jsonl"
    lines = [
        json.dumps({"note_id": "n1", "draft": "旧", "input_sig": "keep|S001|I01"}),
        json.dumps({"note_id": "n2", "status": "llm_failed", "error": "x"}),
        json.dumps({"note_id": "n1", "draft": "新", "input_sig": "keep|S002|I01"}),
        json.dumps({"note_id": "", "draft": "无id"}),
        "坏行",
    ]
    p.write_text("\n".join(lines), encoding="utf-8")
    d = load_drafts(p)
    assert d["n1"]["draft"] == "新" and d["n1"]["input_sig"] == "keep|S002|I01"
    assert "n2" not in d and len(d) == 1


def test_clean_context_strips_noise_keeps_original() -> None:
    note = Note(note_id="n1",
                title="考研失眠怎么办",
                desc=("uu们救命！！有没有姐妹考研期间也天天失眠的，大家都是怎么熬过来的[哭惹R][笑哭R] "
                      "@小助手 真的快撑不住了😭\n\n\n二编：谢谢大家，评论区好温暖 #考研 #失眠怎么办"),
                tags=("考研",), raw={})
    out = clean_context(note)
    assert "考研失眠怎么办" in out                     # 标题并入（正文没重复它）
    assert "天天失眠的，大家都是怎么熬过来的" in out    # 原文措辞原样保留（不概括）
    assert "快撑不住了" in out
    assert "#" not in out and "考研 " != out[:3]        # 话题标签删掉
    assert "[哭惹R]" not in out and "😭" not in out     # 表情码/emoji 删掉
    assert "@" not in out
    assert "二编" not in out and "温暖" not in out      # 二编起的追加段整段砍掉
    assert "!!" not in out and "！！" not in out        # 连排标点压缩
    assert "\n\n" not in out                           # 连续空行压一行


def test_clean_context_title_dedupe_and_no_edit_mark() -> None:
    note = Note(note_id="n2", title="副业求带",
                desc="副业求带，可付费，预算3000，零经验想起步",
                tags=(), raw={})
    assert clean_context(note).count("副业求带") == 1     # 正文已含标题不重复拼
    note2 = Note(note_id="n3", title="", desc="正文只有一段，无标签无表情",
                 tags=(), raw={})
    assert clean_context(note2) == "正文只有一段，无标签无表情"


def test_split_context_query() -> None:
    # 正常 JSON
    ctx, q = split_context_query('{"context": "本科双非机械，绩点 2.8。", '
                                 '"query": "二战跨考计算机上岸可能性？"}')
    assert ctx == "本科双非机械，绩点 2.8。" and q == "二战跨考计算机上岸可能性？"
    # 围栏 + 前后杂文字
    ctx, q = split_context_query('```json\n{"context": "C", "query": "Q"}\n``` 好的')
    assert ctx == "C" and q == "Q"
    # context 缺省 = 空字符串
    ctx, q = split_context_query('{"query": "只有问句"}')
    assert ctx == "" and q == "只有问句"
    # JSON 坏了 → 整段当 query（rewrite_v1 行为）
    ctx, q = split_context_query("  直接一句query  ")
    assert ctx == "" and q == "直接一句query"
    # JSON 有但 query 空 → 也退回整段
    ctx, q = split_context_query('{"context": "只有背景"}')
    assert ctx == "" and q.startswith("{")            # 原样清洗为 query


def test_rewrite_record_roundtrip() -> None:
    r = RewriteRecord(note_id="n1", final_status="keep", scenarios=["S001"],
                      intents=["I01"], rewrite="问句", context="背景",
                      based_on="WH")
    assert rewrite_from_dict(rewrite_to_dict(r)) == r


def test_export_rewrites(tmp_path: Path, taxonomy) -> None:
    notes_path, annotations = _setup(tmp_path)
    task = build_task(annotations, notes_path)
    records = {
        "n1": {"final_status": "keep", "scenarios": ["S001"], "intents": ["I01"],
               "persona": [], "missing_info": "", "comment": "", "rewrite": "大三纠结考研还是就业",
               "context": "大三，机械专业", "based_on": "WH+zlx", "operator": "OP",
               "updated_at": "2026-09-28"},
        "n2": {"final_status": "drop", "scenarios": [], "intents": [], "rewrite": "",
               "based_on": "WH+zlx", "operator": "OP", "updated_at": "2026-09-28"},
    }
    out = tmp_path / "最终改写.xlsx"
    _, st = export_rewrites(task, taxonomy, records, out)
    assert st == {"total": 2, "done": 2, "drop": 1, "rewritten": 1, "todo": 0}

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["最终改写", "待改写", "统计"]
    rows = {r[0]: r for r in wb["最终改写"].iter_rows(min_row=2, values_only=True)}
    assert rows["n1"][3] == "大三纠结考研还是就业"   # 改写query 列
    assert rows["n1"][4] == "大三，机械专业"          # 改写context 列
    assert rows["n1"][5] == "保留" and rows["n2"][5] == "不需要"
    assert rows["n1"][11] == "OP"                    # 操作人列
    assert wb["待改写"].max_row == 1                  # 都处理完了

    # 保留但没改写 → 进待改写，不算完成
    records["n1"]["rewrite"] = ""
    _, st2 = export_rewrites(task, taxonomy, records, out)
    assert st2["rewritten"] == 0 and st2["todo"] == 1 and st2["drop"] == 1


def test_build_rewrite_prompt() -> None:
    note = Note(note_id="n1", title="考研还是就业", desc="大三很纠结" * 300,
                tags=("考研",), raw={})
    inp = RewriteInput(note=note, scenario_paths=("升学 > 考研 > 是否考研",),
                       intent_names=("评估决策",))
    system, user = build_rewrite_prompt("rewrite_v1", inp)
    assert "标准用户查询" in system and "不超过 60 字" in system
    assert "考研还是就业" in user
    assert "升学 > 考研 > 是否考研" in user and "评估决策" in user
    assert "…(截断)" in user                       # 长正文截断
    # v2：context + query 双段 JSON 输出（context=原帖清洗版，不概括）
    system2, user2 = build_rewrite_prompt("rewrite_v2", inp)
    assert '"context"' in system2 and '"query"' in system2
    assert "不概括" in system2 and "JSON" in user2
    # v3：单一明确问题；「大家是怎么处理的」要转成「我该怎么处理」
    system3, _ = build_rewrite_prompt("rewrite_v3", inp)
    assert "明确" in system3 and "大家是怎么处理的」→「我该怎么处理" in system3
    with pytest.raises(KeyError, match="未注册的改写 prompt"):
        build_rewrite_prompt("nope", inp)
