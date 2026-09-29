"""测试夹具：内存构建迷你标签 xlsx 与 Taxonomy。"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from labeler.taxonomy import Scenario, Taxonomy, load_taxonomy


def build_taxonomy_xlsx(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "用户画像"
    ws.append(["字段组", "画像字段", "敏感级别"])
    ws.append(["基础信息", "年龄段", "一般"])
    ws.append(["教育背景", "年级", "一般"])
    ws.append(["教育背景", "学校名称", "一般"])
    ws.append(["职业背景", "求职阶段", "一般"])

    ws2 = wb.create_sheet("场景标签")
    ws2.append(["场景ID", "一级场景", "二级主题", "三级叶节点", "完整路径", "叶节点定义", "建议业务线", "状态"])
    for sid, l1, l2, l3, definition in [
        ("S001", "升学", "考研", "是否考研", "考研与就业等路径选择"),
        ("S002", "升学", "考研", "院校选择", "考研目标院校定位"),
        ("S003", "就业", "国央企", "网申", "央企网申流程"),
        ("S004", "个人成长与生活", "执行与习惯", "拖延", "拖延与启动困难"),
    ]:
        ws2.append([sid, l1, l2, l3, f"{l1} > {l2} > {l3}", definition, "", ""])

    ws3 = wb.create_sheet("意图标签")
    ws3.append(["意图ID", "意图名称", "定义", "典型表达"])
    ws3.append(["I01", "信息获取", "获得事实与资源", "是什么；多少钱"])
    ws3.append(["I02", "评估决策", "判断可行性或在选项间抉择", "A还是B；该不该"])
    ws3.append(["I03", "情绪支持", "被理解安慰共鸣", "求安慰；想倾诉"])

    wb.save(path)


@pytest.fixture
def taxonomy(tmp_path: Path) -> Taxonomy:
    xlsx = tmp_path / "tax.xlsx"
    build_taxonomy_xlsx(xlsx)
    return load_taxonomy(xlsx, additions_path=None)   # 测试用迷你原表，不合并真实补充叶


@pytest.fixture
def scenarios() -> list[Scenario]:
    return [
        Scenario("S001", "升学", "考研", "是否考研", "升学 > 考研 > 是否考研", "考研与就业路径", ""),
        Scenario("S002", "升学", "考研", "院校选择", "升学 > 考研 > 院校选择", "院校定位", ""),
    ]
